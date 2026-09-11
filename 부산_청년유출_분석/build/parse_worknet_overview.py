"""고용24 신규구인 개괄(1-2, 구인업체 소재지 기준) → 시도 × 학력/임금대 tidy.

직종 축과 달리 이 시트에는 **학력별·임금대별** 축이 있다.
부산 기업이 어떤 학력을 요구하고 얼마를 제시하는지 다른 시도와 직접 비교할 수 있다.
"""
import glob, os, re
import openpyxl, pandas as pd

def parse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["신규구인_개괄"]
    rows = list(ws.iter_rows(values_only=True))
    txt = [" ".join(str(c).strip() for c in r if c) if r else "" for r in rows]
    s = next(i for i, t in enumerate(txt) if t.startswith("1-2."))   # 구인업체 소재지 기준
    hi = next(i for i in range(s, len(rows))
              if rows[i] and any(str(c).strip() == "부산광역시" for c in rows[i] if c))
    hdr = [str(c).strip() if c else "" for c in rows[hi]]
    regions = {n: j for j, n in enumerate(hdr) if n.endswith(("시", "도", "특별시", "광역시", "자치도")) or n == "전 체"}
    out, axis = [], None
    for r in rows[hi + 1:]:
        if not r: continue
        a = str(r[1]).strip() if r[1] else ""
        b = str(r[2]).strip() if r[2] else ""
        if a and a.endswith("별"): axis = a.replace("별", "")
        # 항목명은 a 또는 b 중 축 이름이 아닌 쪽
        item = b if b else (a if a and not a.endswith("별") else "")
        if not item or axis is None: continue
        for reg, j in regions.items():
            try: v = int(r[j])
            except (TypeError, ValueError): continue
            out.append({"축": axis, "항목": item, "시도": reg.replace("특별시", "").replace("광역시", ""), "값": v})
        if axis == "임금대" and "250만원 이상" in item: break
    wb.close()
    return pd.DataFrame(out)

frames = []
for f in sorted(glob.glob("data/고용24/*.xlsx")):
    d = parse(f); d.insert(0, "연월", os.path.basename(f)[:7]); frames.append(d)
df = pd.concat(frames, ignore_index=True)
df.to_csv("out/고용24_시도별_학력임금.csv", index=False, encoding="utf-8-sig")
print(f"행 {len(df):,} · 월 {df.연월.nunique()} · 축 {sorted(df.축.unique())} · 시도 {df.시도.nunique()}")

# ── 산업 축 (7-2, 구인업체 소재지 기준)
def parse_ind(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["산업별_구인"]
    rows = list(ws.iter_rows(values_only=True))
    txt = [" ".join(str(c).strip() for c in r if c) if r else "" for r in rows]
    s = next(i for i, t in enumerate(txt) if t.startswith("7-2."))
    hi = next(i for i in range(s, len(rows))
              if rows[i] and any(str(c).strip() == "부산광역시" for c in rows[i] if c))
    hdr = [str(c).strip() if c else "" for c in rows[hi]]
    ci = {n: j for j, n in enumerate(hdr) if n}
    out = []
    for r in rows[hi + 1:]:
        if not r: continue
        ind = str(r[ci["산업대분류"]]).strip() if r[ci["산업대분류"]] else ""
        if not ind or ind == "None": continue
        for reg, j in ci.items():
            if reg == "산업대분류": continue
            try: v = int(r[j])
            except (TypeError, ValueError): continue
            out.append({"산업": ind, "시도": reg.replace("특별시", "").replace("광역시", ""), "값": v})
    wb.close()
    return pd.DataFrame(out)

fi = []
for f in sorted(glob.glob("data/고용24/*.xlsx")):
    d = parse_ind(f); d.insert(0, "연월", os.path.basename(f)[:7]); fi.append(d)
ind = pd.concat(fi, ignore_index=True)
ind.to_csv("out/고용24_시도별_산업구인.csv", index=False, encoding="utf-8-sig")
print(f"산업 축: 행 {len(ind):,} · 산업 {ind.산업.nunique()}종")
